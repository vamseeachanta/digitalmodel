# # Copyright (c) 2024 Digital Model Project
# # Licensed under the MIT License. See LICENSE file for details.

"""Validated fatigue assessment and report export on synthetic loads."""

import json
import re
from pathlib import Path

import numpy as np
import pytest
from jsonschema import Draft202012Validator, ValidationError as SchemaValidationError
from openpyxl import load_workbook
from pydantic import ValidationError

from digitalmodel.fatigue.assessment import AssessmentInput, assess_fatigue
from digitalmodel.fatigue.assessment_export import export_assessment
from digitalmodel.fatigue.sn_curves import get_sn_curve
from digitalmodel.reporting import DocumentMeta, RevisionRow


def histogram_input(**changes):
    """Synthetic block: one quarter of the endurance at a supplied range."""
    curve = get_sn_curve("D")
    endurance = float(curve.cycles(np.array([100.0]))[0])
    data = dict(
        design_years=4.0,
        block_years=2.0,
        design_fatigue_factor=2.0,
        load_cases=[dict(name="synthetic", kind="histogram",
                         stress_ranges_mpa=[100.0], cycle_counts=[endurance / 4])],
    )
    data.update(changes)
    return data


def test_duration_and_design_factor_are_explicit():
    result = assess_fatigue(AssessmentInput.model_validate(histogram_input()))
    assert result.block_damage == pytest.approx(0.25)
    assert result.life_years == pytest.approx(8.0)
    assert result.utilisation == pytest.approx(1.0)
    assert result.passed
    assert result.governing_load_case == "synthetic"
    assert result.cases[0].cycles[0].damage == pytest.approx(0.25)


def test_governing_case_is_largest_damage_contributor():
    data = histogram_input()
    smaller = dict(data["load_cases"][0], name="smaller", cycle_counts=[1.0])
    data["load_cases"].insert(0, smaller)
    result = assess_fatigue(AssessmentInput.model_validate(data))
    assert result.governing_load_case == "synthetic"
    assert not result.passed


def test_constant_history_has_zero_damage_and_unknown_finite_life():
    data = histogram_input(load_cases=[dict(name="constant", kind="history",
                                          stress_history_mpa=[5.0, 5.0, 5.0])])
    result = assess_fatigue(AssessmentInput.model_validate(data))
    assert result.block_damage == 0.0
    assert result.life_years is None
    assert result.passed
    json.dumps(result.model_dump(), allow_nan=False)


def test_calm_history_preserves_storm_case_damage():
    storm = dict(name="storm", kind="history", stress_history_mpa=[0, 100, 0])
    calm = dict(name="calm", kind="history", stress_history_mpa=[10, 10])
    storm_only = assess_fatigue(AssessmentInput.model_validate(
        histogram_input(load_cases=[storm])))
    combined = assess_fatigue(AssessmentInput.model_validate(
        histogram_input(load_cases=[storm, calm])))
    assert combined.block_damage == pytest.approx(storm_only.block_damage)
    assert combined.cases[1].block_damage == 0.0


@pytest.mark.parametrize("changes", [
    {"block_years": 0}, {"design_years": -1}, {"design_fatigue_factor": 0.5},
    {"design_years": float("inf")}, {"environment": "unknown"},
    {"sn_curve_name": "unknown"}, {"design_code": "API"},
    {"unexpected": True}, {"load_cases": []},
    {"load_cases": [dict(name="bad", kind="histogram", stress_ranges_mpa=[-1],
                         cycle_counts=[1])]},
    {"load_cases": [dict(name="bad", kind="histogram", stress_ranges_mpa=[1],
                         cycle_counts=[])]},
    {"load_cases": [dict(name="bad", kind="history",
                         stress_history_mpa=[0, float("nan")])]},
])
def test_invalid_input_fails_before_computation(changes):
    with pytest.raises(ValidationError):
        AssessmentInput.model_validate(histogram_input(**changes))


def test_duplicate_case_names_are_rejected():
    data = histogram_input()
    data["load_cases"] *= 2
    with pytest.raises(ValidationError, match="unique"):
        AssessmentInput.model_validate(data)


def test_schema_is_generated_from_runtime_contract():
    schema = json.loads(Path("docs/skills/fatigue-assessment/input.schema.json").read_text())
    assert schema == AssessmentInput.model_json_schema()
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    validator.validate(histogram_input())
    with pytest.raises(SchemaValidationError, match="minimum"):
        validator.validate(histogram_input(block_years=0))


def test_export_reads_native_output_and_workbook_back(tmp_path):
    inp = AssessmentInput.model_validate(histogram_input())
    document = DocumentMeta(number="B0000-RPT-001-01", revision="01",
                            title="Fatigue analysis", project="Synthetic analysis",
                            client="Synthetic input", date="2026-10-14",
                            revision_history=[
                                RevisionRow(rev="00", date="2026-10-14", description="Issued"),
                                RevisionRow(rev="01", date="2026-10-14",
                                            description="Display precision aligned with report format"),
                            ])
    artifacts = export_assessment(inp, tmp_path, document=document, pdf="auto")
    native = json.loads(artifacts.native_path.read_text())
    assert native["result"]["block_damage"] == pytest.approx(0.25)
    assert native["result"]["utilisation"] == pytest.approx(1.0)
    assert artifacts.report.html_path.is_file()
    html = artifacts.report.html_path.read_text()
    assert html.index('id="summary"') < html.index('id="revision-history"')
    assert "Governing load case" in html
    assert "synthetic" in html
    echo = html.split('id="input-echo"', 1)[1].split("</section>", 1)[0]
    assert not re.search(r"\d+\.\d{2,}", echo)
    book = load_workbook(artifacts.excel_path, data_only=True)
    assert book["Summary"]["B2"].value == pytest.approx(0.25)
    assert book["Cycles"]["E2"].value == pytest.approx(0.25)
    assert book["Summary"]["B2"].number_format == "0.0E+00"
    assert book["Cycles"]["E2"].number_format == "0.0E+00"
    for row in book["Input"].iter_rows(min_row=2):
        assert not re.search(r"\d+\.\d{2,}", row[1].value)
    manifest = json.loads(artifacts.manifest_path.read_text())
    assert manifest["native_sha256"]
    assert manifest["pdf_rendered"] == artifacts.report.pdf_status.rendered
    book.close()


def test_workbook_preserves_supplied_case_names_as_literal_text(tmp_path):
    data = histogram_input()
    data["load_cases"][0]["name"] = "=1+1"
    document = DocumentMeta(number="B0000-RPT-001-00", revision="00",
                            title="Fatigue analysis", project="Synthetic analysis",
                            client="Synthetic input")
    artifacts = export_assessment(AssessmentInput.model_validate(data), tmp_path,
                                  document=document, pdf="auto")
    book = load_workbook(artifacts.excel_path, data_only=False)
    assert book["Summary"]["B5"].value == "=1+1"
    assert book["Summary"]["B5"].data_type == "s"
    assert book["Cycles"]["A2"].data_type == "s"
    book.close()


def test_native_readback_bytes_are_checked_before_report_export(tmp_path, monkeypatch):
    original = Path.read_bytes

    def tampered_read(path):
        if path.name == "fatigue-native.json":
            return b"{}"
        return original(path)

    monkeypatch.setattr(Path, "read_bytes", tampered_read)
    document = DocumentMeta(number="B0000-RPT-001-00", revision="00",
                            title="Fatigue analysis", project="Synthetic analysis",
                            client="Synthetic input")
    with pytest.raises(ValueError, match="readback differs"):
        export_assessment(AssessmentInput.model_validate(histogram_input()), tmp_path,
                          document=document, pdf="auto")
    assert not (tmp_path / "fatigue-analysis.html").exists()
