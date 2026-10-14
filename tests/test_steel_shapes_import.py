"""Offline workbook import tests; hand-written synthetic profiles only."""

import hashlib
import json
import socket
from importlib.resources import files

import pytest

from digitalmodel.steel_shapes import get_section_data
from digitalmodel.steel_shapes_import import import_workbook


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def denied(*args, **kwargs):
        raise AssertionError("workbook import tests shall remain offline")

    monkeypatch.setattr(socket.socket, "connect", denied)


def _dimension_fields(item):
    dims = item["dimensions"]
    family = item["family"]
    fields = {"A": item["area"], "kdes": dims.get("kdes"), "tnom": dims.get("tnom")}
    if family in {"W", "S", "HP", "C"}:
        fields.update({k: dims[k] for k in ("d", "bf", "tw", "tf")})
    elif family == "L":
        fields.update({"d": dims["d"], "b": dims["b"], "t": dims["t"]})
    elif family == "HSS_RECT":
        fields.update({"Ht": dims["d"], "B": dims["b"], "tdes": dims["t"]})
    else:
        fields.update({"OD": dims["d"], "tdes": dims["t"]})
    return fields


@pytest.fixture
def workbook(tmp_path):
    from openpyxl import Workbook

    data = json.loads(
        files("digitalmodel.data_systems.data")
        .joinpath("steel_shapes_sample.json")
        .read_text()
    )
    book = Workbook()
    sheet = book.active
    sheet.title = "Database v16.0"
    # Same duplicate headings as the source: imperial block followed by SI.
    headings = [
        "Type",
        "EDI_Std_Nomenclature",
        "AISC_Manual_Label",
        "A",
        "d",
        "Ht",
        "OD",
        "bf",
        "B",
        "b",
        "tw",
        "tf",
        "t",
        "tnom",
        "tdes",
        "kdes",
        "EDI_Std_Nomenclature",
        "AISC_Manual_Label",
    ]
    sheet.append(headings)
    for item in data["sections"]:
        family = item["family"]
        fields = _dimension_fields(item)
        row = [fields.get(h) for h in headings]
        row[:3] = [
            "HSS" if family.startswith("HSS") else family,
            item["aliases"][0],
            item["name"],
        ]
        row[-2:] = [item["metric_name"], item["metric_name"]]
        sheet.append(row)
        sheet.cell(
            row=sheet.max_row, column=headings.index("A") + 1
        ).number_format = item["source"].get("area_number_format", "0.00")
    path = tmp_path / "source.xlsx"
    book.save(path)
    return path, data


def test_import_preserves_all_families_and_aliases(workbook, tmp_path):
    path, expected = workbook
    data = import_workbook(path, source_url="https://www.aisc.org/example.xlsx")
    assert len(data["sections"]) == len(expected["sections"])
    assert data["source"]["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    output = tmp_path / "catalog.json"
    output.write_text(json.dumps(data))
    for item in expected["sections"]:
        actual = get_section_data(item["metric_name"], catalog_path=output)
        assert actual["dimensions"] == item["dimensions"]
        # XLSX numeric serialization retains about 15 significant decimal digits.
        assert actual["area"] == pytest.approx(item["area"], rel=1e-14)
        assert (
            actual["source"]["area_number_format"]
            == item["source"].get("area_number_format", "0.00")
        )


def test_reject_wrong_workbook_edition(workbook):
    from openpyxl import load_workbook

    path, _ = workbook
    book = load_workbook(path)
    book.active.title = "wrong edition"
    book.save(path)
    with pytest.raises(ValueError, match="Database v16.0"):
        import_workbook(path, source_url="https://www.aisc.org/example.xlsx")


def test_reject_missing_metric_heading(workbook):
    from openpyxl import load_workbook

    path, _ = workbook
    book = load_workbook(path)
    book.active.cell(row=1, column=book.active.max_column).value = "missing"
    book.save(path)
    with pytest.raises(ValueError, match="heading"):
        import_workbook(path, source_url="https://www.aisc.org/example.xlsx")


def test_reject_missing_required_shape_value(workbook):
    from openpyxl import load_workbook

    path, _ = workbook
    book = load_workbook(path)
    headings = [c.value for c in book.active[1]]
    book.active.cell(row=2, column=headings.index("tw") + 1).value = None
    book.save(path)
    with pytest.raises(ValueError, match="tw"):
        import_workbook(path, source_url="https://www.aisc.org/example.xlsx", strict=True)


def _append_bad_rows(path):
    from openpyxl import load_workbook
    book = load_workbook(path)
    sheet = book.active
    headings = [c.value for c in sheet[1]]
    good = [c.value for c in sheet[2]]
    cases = [("SYN-BAD-GEOMETRY", "tf", 100), ("SYN-MISSING", "tw", None)]
    numbers = []
    for name, field, value in cases:
        row = good.copy()
        row[1:3] = [name, name]
        row[-2:] = ["SYN-METRIC-" + name, "SYN-METRIC-" + name]
        row[headings.index(field)] = value
        sheet.append(row)
        numbers.append(sheet.max_row)
    sheet.append(good)
    numbers.append(sheet.max_row)
    row = good.copy()
    row[2] = "SYN-ALIAS-COLLISION"
    row[-2:] = ["SYN-METRIC-NEW", "SYN-METRIC-NEW"]
    sheet.append(row)
    numbers.append(sheet.max_row)
    book.save(path)
    book.close()
    return numbers


def test_partial_import_reports_each_bad_row(workbook, capsys):
    path, expected = workbook
    numbers = _append_bad_rows(path)
    data = import_workbook(path, source_url="https://www.aisc.org/example.xlsx")
    report = data["import_report"]
    assert report["imported"] == len(expected["sections"])
    assert report["rejected"] == 4
    assert [r["row"] for r in report["rejections"]] == numbers
    reasons = [r["reason"] for r in report["rejections"]]
    assert "geometry" in reasons[0]
    assert "tw" in reasons[1]
    assert "Duplicate" in reasons[2] and "Duplicate" in reasons[3]
    output = capsys.readouterr().err
    for row in report["rejections"]:
        assert str(row["row"]) in output
        assert row["designation"] in output
        assert row["reason"] in output
    assert get_section_data("SYN-W")["synthetic"] is True


def test_strict_import_is_all_or_nothing(workbook, capsys):
    path, _ = workbook
    _append_bad_rows(path)
    with pytest.raises(ValueError, match="Strict import rejected 4 row"):
        import_workbook(path, source_url="https://www.aisc.org/example.xlsx", strict=True)
    assert "SYN-MISSING" in capsys.readouterr().err


def test_no_valid_rows_fails(workbook, capsys):
    from openpyxl import load_workbook
    path, _ = workbook
    book = load_workbook(path)
    sheet = book.active
    sheet.delete_rows(3, sheet.max_row)
    headings = [c.value for c in sheet[1]]
    sheet.cell(row=2, column=headings.index("tw") + 1).value = None
    book.save(path)
    book.close()
    with pytest.raises(ValueError, match="No valid"):
        import_workbook(path, source_url="https://www.aisc.org/example.xlsx")
    assert "row 2" in capsys.readouterr().err


def test_tapered_geometry_rejected_per_row(workbook):
    from openpyxl import load_workbook
    path, expected = workbook
    book = load_workbook(path)
    headings = [c.value for c in book.active[1]]
    # Basic dimension checks pass, but the tangency radius would be negative.
    book.active.cell(row=4, column=headings.index("kdes") + 1).value = .8
    book.save(path)
    book.close()
    data = import_workbook(path, source_url="https://www.aisc.org/example.xlsx")
    assert data["import_report"]["imported"] == len(expected["sections"]) - 1
    assert data["import_report"]["rejections"][0]["row"] == 4
    assert "Tapered" in data["import_report"]["rejections"][0]["reason"]


def test_rejected_aliases_do_not_poison_later_rows(workbook):
    from openpyxl import load_workbook
    path, expected = workbook
    _append_bad_rows(path)
    book = load_workbook(path)
    sheet = book.active
    later = [cell.value for cell in sheet[2]]
    later[1:3] = ["SYN-LATER-ALIAS", "SYN-ALIAS-COLLISION"]
    later[-2:] = ["SYN-METRIC-NEW", "SYN-METRIC-NEW"]
    sheet.append(later)
    book.save(path)
    book.close()
    data = import_workbook(path, source_url="https://www.aisc.org/example.xlsx")
    assert data["import_report"]["imported"] == len(expected["sections"]) + 1
    assert data["import_report"]["rejected"] == 4
    assert data["sections"][-1]["name"] == "SYN-ALIAS-COLLISION"
    assert data["sections"][-1]["metric_name"] == "SYN-METRIC-NEW"


def test_unsupported_families_counted_as_skipped(workbook):
    from openpyxl import load_workbook
    path, expected = workbook
    book = load_workbook(path)
    row = [cell.value for cell in book.active[2]]
    row[0] = "UNSUPPORTED"
    book.active.append(row)
    book.save(path)
    book.close()
    data = import_workbook(path, source_url="https://www.aisc.org/example.xlsx")
    assert data["import_report"]["imported"] == len(expected["sections"])
    assert data["import_report"]["rejected"] == 0
    assert data["import_report"]["skipped"] == 1
