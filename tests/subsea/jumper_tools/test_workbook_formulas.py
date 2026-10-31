# Copyright (c) 2024 Digital Model Project
# Licensed under the MIT License. See LICENSE file for details.
"""Synthetic workbooks test extraction, never spreadsheet calculation parity."""

import hashlib
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock
from zipfile import ZipFile
from xml.etree import ElementTree

import openpyxl
import pytest

from digitalmodel.subsea.jumper_tools.workbook_formulas import extract_formulas


def test_extract_formulas_without_evaluation_or_mutation(tmp_path):
    """Two formulas are written literally; openpyxl supplies no cached results."""
    path = tmp_path / "synthetic.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Inputs"
    wb.active.append([2, 3, "=A1+B1"])
    wb.create_sheet("Geometry")["B2"] = "=Inputs!A1*2"
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    result = extract_formulas(path)
    assert result.source_sha256 == hashlib.sha256(before).hexdigest()
    assert [(f.sheet, f.cell, f.formula, f.cached_value) for f in result.formulas] == [
        ("Inputs", "C1", "=A1+B1", None),
        ("Geometry", "B2", "=Inputs!A1*2", None),
    ]
    assert result.comparison_status == "NOT_EVALUATED"
    assert path.read_bytes() == before


def test_no_formulas_is_explicit(tmp_path):
    path = tmp_path / "values.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = 5
    wb.save(path)
    wb.close()
    assert extract_formulas(path).formulas == ()


def test_missing_workbook():
    with pytest.raises(FileNotFoundError):
        extract_formulas("unavailable.xlsx")


def test_cached_result_and_hidden_sheet_remain_unverified(tmp_path):
    """Formula 2+3 independently gives 5, but the stored cache is deliberately 99.

    Extraction preserves the stale value and NOT_EVALUATED, rather than treating
    either a cached value or a hidden sheet as evidence of calculation parity.
    """
    path = tmp_path / "cached.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.create_sheet("Hidden")
    sheet.sheet_state = "hidden"
    sheet["A1"] = "=2+3"
    wb.save(path)
    wb.close()
    snapshot = BytesIO(path.read_bytes())
    with ZipFile(snapshot) as original, ZipFile(path, "w") as updated:
        for member in original.infolist():
            payload = original.read(member.filename)
            if member.filename == "xl/worksheets/sheet2.xml":
                xml = ElementTree.fromstring(payload)
                cache = xml.find(".//{*}v")
                assert cache is not None
                cache.text = "99"
                payload = ElementTree.tostring(xml)
            updated.writestr(member, payload)
    record = extract_formulas(path)
    assert record.formulas[0].sheet == "Hidden"
    assert record.formulas[0].formula == "=2+3"
    assert record.formulas[0].cached_value == 99
    assert record.comparison_status == "NOT_EVALUATED"


def test_unsupported_array_formula_fails_explicitly(tmp_path):
    from openpyxl.worksheet.formula import ArrayFormula

    path = tmp_path / "array.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = ArrayFormula(ref="A1:A2", text="=ROW(A1:A2)")
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="ordinary text formulas"):
        extract_formulas(path)


@pytest.mark.parametrize("dimension", ["sheet", "row", "cell"])
@pytest.mark.parametrize("longer_view", ["formulas", "cached"])
def test_mismatched_views_fail_explicitly_and_close_workbooks(
    tmp_path, monkeypatch, dimension, longer_view
):
    """Neither view may silently lose an extra sheet, row or cell."""
    path = tmp_path / "mismatched.xlsx"
    path.write_bytes(b"mocked workbook snapshot")
    formula_cell = SimpleNamespace(data_type="f", value="=2+3", coordinate="A1")
    shapes = {
        "formulas": [[(formula_cell,)]],
        "cached": [[(SimpleNamespace(value=5),)]],
    }
    longer = shapes[longer_view]
    if dimension == "sheet":
        longer.append(longer[0])
    elif dimension == "row":
        longer[0].append(longer[0][0])
    else:
        longer[0][0] += (longer[0][0][0],)
    workbooks = [
        Mock(
            worksheets=[
                Mock(title="Synthetic", iter_rows=Mock(return_value=iter(rows)))
                for rows in sheets
            ]
        )
        for sheets in shapes.values()
    ]
    monkeypatch.setattr(openpyxl, "load_workbook", Mock(side_effect=workbooks))

    with pytest.raises(ValueError, match="zip"):
        extract_formulas(path)

    for workbook in workbooks:
        workbook.close.assert_called_once_with()
