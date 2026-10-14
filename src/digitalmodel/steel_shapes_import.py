"""Offline AISC v16 workbook conversion; network access belongs to the CLI.

The repository ships a synthetic sample for tests only. For real section
properties, obtain the shapes database from the publisher yourself, under the
publisher's terms, then run the importer on your local copy; nothing imported
is committed. Publisher: https://www.aisc.org/aisc/publications/steel-construction-manual/aisc-shapes-database-v160/
"""

from __future__ import annotations

import datetime as dt
import hashlib
import sys
from io import BytesIO
from pathlib import Path

from digitalmodel.steel_shapes import validate_catalog

_FAMILY_COLUMNS = {
    "W": {"d": "d", "bf": "bf", "tw": "tw", "tf": "tf", "kdes": "kdes"},
    "HP": {"d": "d", "bf": "bf", "tw": "tw", "tf": "tf", "kdes": "kdes"},
    "S": {"d": "d", "bf": "bf", "tw": "tw", "tf": "tf", "kdes": "kdes"},
    "C": {"d": "d", "bf": "bf", "tw": "tw", "tf": "tf", "kdes": "kdes"},
    "L": {"d": "d", "b": "b", "t": "t", "kdes": "kdes"},
    "HSS_RECT": {"d": "Ht", "b": "B", "t": "tdes", "tnom": "tnom"},
    "HSS_ROUND": {"d": "OD", "t": "tdes", "tnom": "tnom"},
    "PIPE": {"d": "OD", "t": "tdes", "tnom": "tnom"},
}


def _header_indices(header: tuple) -> tuple[dict, dict]:
    """Split duplicate imperial/SI headings without hardcoded column offsets."""
    edi = [i for i, value in enumerate(header) if value == "EDI_Std_Nomenclature"]
    manual = [i for i, value in enumerate(header) if value == "AISC_Manual_Label"]
    if len(edi) != 2 or len(manual) != 2 or not edi[0] < manual[0] < edi[1] < manual[1]:
        raise ValueError("Missing/ambiguous imperial and metric name heading blocks")
    imperial = {
        value: i for i, value in enumerate(header[: edi[1]]) if value is not None
    }
    metric = {"EDI_Std_Nomenclature": edi[1], "AISC_Manual_Label": manual[1]}
    required = {
        "Type",
        "A",
        "d",
        "Ht",
        "OD",
        "bf",
        "B",
        "b",
        "tw",
        "tf",
        "t",
        "tnom",
        "tdes",
        "kdes",
    }
    missing = required - imperial.keys()
    if missing:
        raise ValueError(f"Missing AISC dimension heading(s): {sorted(missing)}")
    return imperial, metric


def _record(row: tuple, number: int, imperial: dict, metric: dict) -> dict | None:
    from openpyxl.utils import get_column_letter

    family = row[imperial["Type"]]
    if family == "HSS":
        diameter = row[imperial["OD"]]
        family = "HSS_ROUND" if isinstance(diameter, (int, float)) else "HSS_RECT"
    if family not in _FAMILY_COLUMNS:
        return None
    mapping = _FAMILY_COLUMNS[family]
    dims = {key: row[imperial[column]] for key, column in mapping.items()}
    names = {
        "name": row[imperial["AISC_Manual_Label"]],
        "metric_name": row[metric["AISC_Manual_Label"]],
    }
    columns = {
        key: get_column_letter(imperial[column] + 1) for key, column in mapping.items()
    }
    columns.update(
        name=get_column_letter(imperial["AISC_Manual_Label"] + 1),
        metric_name=get_column_letter(metric["AISC_Manual_Label"] + 1),
        area=get_column_letter(imperial["A"] + 1),
    )
    return {
        **names,
        "family": family,
        "dimensions": dims,
        "area": row[imperial["A"]],
        "aliases": [row[imperial["EDI_Std_Nomenclature"]]],
        "metric_aliases": [row[metric["EDI_Std_Nomenclature"]]],
        "source": {"sheet": "Database v16.0", "row": number, "columns": columns},
    }


def _catalog_metadata(raw: bytes, source_url: str) -> dict:
    return {
        "schema": "aisc-shapes/1",
        "units": {"length": "in", "area": "in^2"},
        "source": {
            "catalog_id": "aisc-shapes",
            "publication": "AISC Shapes Database v16.0, August 2023",
            "url": source_url,
            "sha256": hashlib.sha256(raw).hexdigest(),
            "retrieved_on": dt.datetime.now(dt.timezone.utc).date().isoformat(),
            "access_note": "User-supplied workbook; verify source, terms and edition before engineering use.",
        },
        "sections": [],
    }


def _validated_rows(rows, imperial: dict, metric: dict, data: dict) -> dict:
    records, aliases, rejections = [], {}, []
    skipped = 0
    for number, cells in enumerate(rows, start=2):
        row = tuple(cell.value for cell in cells)
        designation = row[imperial["AISC_Manual_Label"]]
        try:
            item = _record(row, number, imperial, metric)
            if item is None:
                skipped += 1
                continue
            item["source"]["area_number_format"] = cells[imperial["A"]].number_format
            index = validate_catalog({**data, "sections": [item]})
            for key in index:
                if key in aliases:
                    raise ValueError(f"Duplicate {key[0]} section alias: {key[1]}")
        except (ValueError, TypeError, IndexError) as exc:
            rejection = {"row": number, "designation": str(designation or "(unreadable)"),
                         "reason": str(exc)}
            rejections.append(rejection)
            print(f"Rejected row {number} ({rejection['designation']}): {exc}", file=sys.stderr)
            continue
        # Commit aliases only after the entire row passes; rejected rows reserve none.
        aliases.update(index)
        records.append(item)
    data["sections"] = records
    return {"imported": len(records), "rejected": len(rejections),
            "skipped": skipped, "rejections": rejections}


def import_workbook(path: str | Path, *, source_url: str, strict: bool = False) -> dict:
    """Convert supported rows independently, preserving source cells and aliases.

    Returns the catalogue dictionary with ``import_report`` counts and individual
    rejections (worksheet row number, readable designation and reason). Diagnostics
    go to stderr. Unsupported families are excluded and counted as skipped.
    ``strict=True`` restores all-or-nothing rejection; zero valid rows also fail.
    No file is written and no network access occurs here. The committed sample
    is synthetic, for tests only; import a user-obtained publisher workbook for
    real properties. Nothing imported is committed.
    """
    from openpyxl import load_workbook

    raw = Path(path).read_bytes()
    data = _catalog_metadata(raw, source_url)
    book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        if "Database v16.0" not in book.sheetnames:
            raise ValueError("Expected AISC worksheet 'Database v16.0'")
        rows = iter(book["Database v16.0"].iter_rows())
        header = next(rows, None)
        if header is None:
            raise ValueError("Missing workbook heading row")
        imperial, metric = _header_indices(tuple(cell.value for cell in header))
        report = _validated_rows(rows, imperial, metric, data)
    finally:
        book.close()
    data["import_report"] = report
    print(f"Imported {report['imported']} rows; rejected {report['rejected']}; "
          f"skipped {report['skipped']} unsupported rows", file=sys.stderr)
    if strict and report["rejected"]:
        reasons = "; ".join(item["reason"] for item in report["rejections"])
        raise ValueError(f"Strict import rejected {report['rejected']} rows: {reasons}")
    if not report["imported"]:
        raise ValueError("No valid supported rows in imported workbook")
    validate_catalog(data)
    return data
