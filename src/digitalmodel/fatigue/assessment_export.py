# # Copyright (c) 2024 Digital Model Project
# # Licensed under the MIT License. See LICENSE file for details.

"""Native-output readback and shared-engine fatigue report export."""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from digitalmodel.fatigue.assessment import AssessmentInput, AssessmentResult, assess_fatigue
from digitalmodel.fatigue.c203_editions import (
    DNV_RP_C203_IMPLEMENTED_EDITION, DNV_RP_C203_VERIFIED_EDITION,
)
from digitalmodel.reporting import (
    DocumentMeta, Provenance, ReportArtifacts, ReportSpec, Section, StandardLabel,
    StatusBlock, TableBlock, TextBlock, write_report,
)
from digitalmodel.reporting.pdf import PdfMode


@dataclass(frozen=True)
class AssessmentArtifacts:
    """Files written from the same read-back native result."""

    native_path: Path
    excel_path: Path
    manifest_path: Path
    report: ReportArtifacts


def _write_json(path: Path, data: dict[str, Any]) -> None:
    """Persist finite canonical JSON with deterministic key order."""
    path.write_text(json.dumps(data, sort_keys=True, indent=2, allow_nan=False) + "\n",
                    encoding="utf-8")


def _summary_rows(result: AssessmentResult) -> list[list[Any]]:
    """Return native summary values without recalculation."""
    return [["Block damage", result.block_damage],
            ["Life (years)", result.life_years],
            ["Design utilisation", result.utilisation],
            ["Governing load case", result.governing_load_case],
            ["Criterion", "Design utilisation <= 1.0"],
            ["Result", "PASS" if result.passed else "FAIL"]]


def _display_values(value: Any) -> Any:
    """Format input echo numbers without modifying the native calculation."""
    if isinstance(value, float):
        return f"{value:.1e}"
    if isinstance(value, list):
        return [_display_values(item) for item in value]
    if isinstance(value, dict):
        return {key: _display_values(item) for key, item in value.items()}
    return value


def _write_excel(path: Path, result: AssessmentResult, inp: AssessmentInput) -> None:
    """Write native result tables and echo the engineering input."""
    book = Workbook()
    summary = book.active
    if summary is None:
        raise RuntimeError("workbook has no summary sheet")
    summary.title = "Summary"
    summary.append(["Quantity", "Value"])
    for row in _summary_rows(result):
        summary.append(row)
    cycles = book.create_sheet("Cycles")
    cycles.append(["Case", "Range (MPa)", "Cycles", "Allowable cycles", "Damage"])
    for case in result.cases:
        for cycle in case.cycles:
            cycles.append([case.name, cycle.stress_range_mpa, cycle.cycles,
                           cycle.allowable_cycles, cycle.damage])
    echo = book.create_sheet("Input")
    echo.append(["Quantity", "Value"])
    for key, value in inp.model_dump().items():
        echo.append([key, json.dumps(_display_values(value), allow_nan=False)])
    for sheet in (summary, cycles):
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.0E+00"
    book.save(path)
    book.close()


def _report_sections(result: AssessmentResult, document: DocumentMeta) -> list[Section]:
    """Bind the reported conclusion to its case and utilisation criterion."""
    summary = [[key, f"{value:.1e}" if isinstance(value, float) else value]
               for key, value in _summary_rows(result)]
    revisions = [[row.rev, row.date, row.description]
                 for row in document.revision_history]
    return [
        Section(key="summary", title="Summary", blocks=[
            TableBlock(title="Fatigue analysis results", columns=["Quantity", "Value"],
                       rows=summary),
            StatusBlock(label="Design fatigue utilisation",
                        status="PASS" if result.passed else "FAIL",
                        governing_case=result.governing_load_case,
                        detail="Acceptance requires design utilisation <= 1.0."),
        ]),
        Section(key="revision-history", title="Revision history", blocks=[
            TableBlock(title="Document revisions", columns=["Revision", "Date", "Change"],
                       rows=revisions),
        ]),
        Section(key="basis", title="Basis and limitations", blocks=[TextBlock(markdown=(
            "Exact counted ranges are passed to the native Miner damage module. "
            "Cases are additive within one declared load block. Life equals block "
            "duration divided by block damage; design utilisation equals block damage "
            "times design/block duration times the supplied design fatigue factor. "
            "Null life means zero damage under the supplied loading, not measured life.\n\n"
            "Displayed numbers use one decimal in scientific notation. Native JSON "
            "retains calculation precision.\n\n"
            f"DNV-RP-C203, Fatigue design of offshore steel structures, "
            f"{DNV_RP_C203_IMPLEMENTED_EDITION}, S-N curves. Library values were checked "
            f"against {DNV_RP_C203_VERIFIED_EDITION}; agreement with the implemented "
            "edition is not established. ASTM E1049-85, Standard Practices for "
            "Cycle Counting in Fatigue Analysis, Rainflow Counting.\n\n"
            "Curve classification, mean-stress corrections, thickness corrections "
            "and load representativeness require an engineering basis. If utilisation "
            "exceeds 1.0, revise loading or the component and repeat the analysis. "
            "Before use, confirm the selected curve, environment and block duration."
        ))]),
    ]


def export_assessment(
    inp: AssessmentInput,
    output_dir: Path | str,
    *,
    document: DocumentMeta,
    pdf: PdfMode = "auto",
) -> AssessmentArtifacts:
    """Calculate, persist, read back, then export HTML/PDF and Excel.

    PDF availability is recorded by the shared report engine. The native
    JSON and workbook remain available if no renderer is installed. All
    output paths are beneath the caller-supplied directory.
    """
    inp = AssessmentInput.model_validate(inp.model_dump())
    result = assess_fatigue(inp)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    native_path = output_dir / "fatigue-native.json"
    payload = {"input": inp.model_dump(), "result": result.model_dump()}
    _write_json(native_path, payload)
    readback = json.loads(native_path.read_text(encoding="utf-8"))
    if readback != payload:
        raise ValueError("native fatigue output readback differs from calculated output")
    result = AssessmentResult.model_validate(readback["result"])
    digest = hashlib.sha256(native_path.read_bytes()).hexdigest()
    provenance = Provenance().add("native_calculation", native_path.name, digest=digest)
    spec = ReportSpec(document=document, sections=_report_sections(result, document),
                      standards=[StandardLabel(code_id=inp.design_code,
                                               edition=DNV_RP_C203_IMPLEMENTED_EDITION,
                                               provenance="Repository S-N library")],
                      provenance=provenance, input_echo=_display_values(readback["input"]))
    report = write_report(spec, output_dir, "fatigue-analysis", pdf=pdf)
    excel_path = output_dir / "fatigue-analysis.xlsx"
    _write_excel(excel_path, result, inp)
    manifest_path = output_dir / "fatigue-export.json"
    _write_json(manifest_path, {"native": native_path.name, "native_sha256": digest,
                               "excel": excel_path.name, "html": report.html_path.name,
                               "pdf_rendered": report.pdf_status.rendered})
    return AssessmentArtifacts(native_path, excel_path, manifest_path, report)
