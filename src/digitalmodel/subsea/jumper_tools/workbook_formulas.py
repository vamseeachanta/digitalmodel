# Copyright (c) 2024 Digital Model Project
# Licensed under the MIT License. See LICENSE file for details.
"""Read workbook formulas and cached values without recalculation or writes.

openpyxl is optional until extract_formulas is called. A cached value can be
absent or stale; its presence never establishes a verified calculation result.
Only trusted, reasonably sized workbook inputs should be supplied.
"""

import hashlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path


@dataclass(frozen=True)
class WorkbookFormula:
    """Formula location and text, with the unverified stored result if present."""

    sheet: str
    cell: str
    formula: str
    cached_value: object


@dataclass(frozen=True)
class WorkbookExtraction:
    """Digest binds extraction to exact input bytes; no source path is exported."""

    source_sha256: str
    formulas: tuple[WorkbookFormula, ...]
    comparison_status: str = "NOT_EVALUATED"


def extract_formulas(path: str | Path) -> WorkbookExtraction:
    """Read ordinary formulas from all sheets, retaining hidden-sheet formulas.

    The same immutable byte snapshot supplies both formula and cached-value
    views. Unsupported non-text array/data-table formulas fail explicitly.
    Extracted sheet names and formula text may contain confidential data;
    extraction does not authorize publication of the resulting records.
    """
    from openpyxl import load_workbook

    payload = Path(path).read_bytes()
    formulas = load_workbook(
        BytesIO(payload), data_only=False, read_only=True, keep_links=False
    )
    try:
        cached = load_workbook(
            BytesIO(payload), data_only=True, read_only=True, keep_links=False
        )
        try:
            records = []
            for sheet, values in zip(formulas.worksheets, cached.worksheets):
                for row, value_row in zip(sheet.iter_rows(), values.iter_rows()):
                    for cell, value in zip(row, value_row):
                        if cell.data_type != "f":
                            continue
                        if not isinstance(cell.value, str):
                            raise ValueError(
                                "only ordinary text formulas are supported"
                            )
                        records.append(
                            WorkbookFormula(
                                sheet.title, cell.coordinate, cell.value, value.value
                            )
                        )
            return WorkbookExtraction(
                hashlib.sha256(payload).hexdigest(), tuple(records)
            )
        finally:
            cached.close()
    finally:
        formulas.close()
