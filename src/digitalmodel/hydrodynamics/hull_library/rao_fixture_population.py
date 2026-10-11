"""Populate the bounded public benchmark fixture database (issue 1736).

Inputs remain the repository spreadsheets; no hull scales or solver qualification
are inferred. General RAODatabase and RaoRegistry storage contracts are unchanged.
"""

import hashlib
import json
import os
import re
from datetime import date
from pathlib import Path
import subprocess

import numpy as np
import openpyxl

from digitalmodel.hydrodynamics.hull_library.rao_database import RAODatabase
from digitalmodel.hydrodynamics.hull_library.rao_extractor import (
    XlsxFormatDetector,
    populate_database_from_xlsx,
)
from digitalmodel.hydrodynamics.hull_library.rao_fixture_comparison import (
    FIELDS,
    compare_box_exports,
)

FIXTURE_STEMS = (
    "test01_unit_box", "ellipsoid", "L00_test01",
    "L01_001_ship_raos", "L02_OC4_semi_sub",
)
BOX_PAIR = {"test01_unit_box", "L00_test01"}
EXPECTED_IDS = set(FIXTURE_STEMS) - {"L00_test01"}
MISSING = "not stated in source"
DOFS = ("surge", "sway", "heave", "roll", "pitch", "yaw")
METADATA_FIELDS = ("units", "conventions", "source", "loading_condition")
# Retain the original declared storage basis for identity, not as source evidence.
BASIS = {
    "frequencies": "rad/s", "directions": "deg", "phases": "deg",
    "translations": "m/m", "rotations": "rad/m",
}


def _labels(sheet, cells):
    return [{"text": cell.value, "sheet": sheet.title, "cell": cell.coordinate}
            for cell in cells if isinstance(cell.value, str)]


def _unit(labels, source_unit, stored_unit=None, conversion="none"):
    return {"labels": labels, "source_unit": source_unit,
            "stored_unit": stored_unit or source_unit, "conversion": conversion}


def _read_units(wb, fmt):
    units = {}
    if fmt == "native":
        sheet = wb["Displacement RAOs"]
        axis = sheet["B9"].value or MISSING
        units["frequencies"] = _unit(
            _labels(sheet, [sheet["B7"], sheet["B9"]]), axis, "rad/s",
            "omega = 2*pi/T" if axis == "s" else "none")
        units["directions"] = _unit(_labels(sheet, [sheet["A9"]]),
                                     sheet["A9"].value or MISSING)
        for dof, col in zip(DOFS, (3, 5, 7, 9, 11, 13)):
            cell = sheet.cell(9, col)
            units[dof] = _unit(_labels(sheet, [cell]), cell.value or MISSING)
        cells = [sheet.cell(9, col) for col in (4, 6, 8, 10, 12, 14)]
        phase_units = {cell.value or MISSING for cell in cells}
        if len(phase_units) != 1:
            raise ValueError("Inconsistent native phase units")
        units["phases"] = _unit(_labels(sheet, cells), phase_units.pop())
        expected = dict(zip(DOFS, ["m/m"] * 3 + ["rad/m"] * 3))
        expected.update(directions="deg", phases="deg")
        for key, value in expected.items():
            if units[key]["source_unit"] not in (value, MISSING):
                raise ValueError(f"Unsupported source unit for {key}")
        return units
    sheet = wb["RAOs"]
    if sheet["A1"].value != "Frequency (rad/s)":
        raise ValueError("Unsupported pipeline frequency label")
    units["frequencies"] = _unit(_labels(sheet, [sheet["A1"]]), "rad/s")
    headers = list(sheet[1])
    heading_labels = _labels(sheet, [c for c in headers
                                    if str(c.value).startswith("Surge_Mag_H")])
    explicit = [c for name in ("Summary", "Discretization") for row in wb[name]
                for c in row if c.value in ("Heading range (deg)", "Heading (deg)")]
    heading_labels += [label for c in explicit for label in _labels(wb[c.parent.title], [c])]
    units["directions"] = _unit(heading_labels, "deg" if explicit else MISSING)
    for dof in DOFS:
        units[dof] = _unit(_labels(sheet, [c for c in headers
                            if str(c.value).startswith(dof.title() + "_Mag_H")]), MISSING)
    units["phases"] = _unit(_labels(sheet, [c for c in headers
                                          if "_Phase_H" in str(c.value)]), MISSING)
    return units


def _statement(sheet, cells, statement=None, values=None):
    result = {"statement": statement or cells[0].value, "sheet": sheet.title,
              "cells": [cell.coordinate for cell in cells if cell.value is not None]}
    if values is not None:
        result["values"] = values
    return result


def _read_source(wb, fmt):
    sheet = wb["Summary"]
    result = dict.fromkeys(("description", "solver", "version", "model_or_run", "date"), MISSING)
    if fmt == "pipeline":
        for row in sheet:
            key = {"Job": "model_or_run", "Solver": "solver"}.get(row[0].value)
            if key:
                result[key] = _statement(sheet, list(row[:2]), row[1].value)
        return result
    cell = sheet["A2"]
    result["description"] = _statement(sheet, [cell])
    match = re.fullmatch(r"(\S+) ([^:]+): (.+) \(modified (.+) by .+\)", str(cell.value))
    if match:
        for key, text in zip(("solver", "version", "model_or_run"), match.groups()[:3]):
            result[key] = _statement(sheet, [cell], text)
        date_text = re.search(r"\d{1,2}/\d{1,2}/\d{4}", match[4])
        if date_text:
            result["date"] = _statement(sheet, [cell], date_text[0])
    return result


def _read_loading(wb):
    names = {
        "draft": ("Draft (m)",), "displacement": ("Displacement (te)",),
        "mass": ("Mass (te)",), "centre_of_gravity": ("Centre of gravity (m)",),
        "centre_of_mass": ("Centre of mass (m)", "Centre of mass Z (m)"),
        "radii_of_gyration": ("Radii of gyration (m)",),
        "water_depth": ("Water depth (m)",), "body_origin": ("Body origin",),
        "mesh_position": ("Mesh position (m)",), "body_position": ("Body position (m)",),
        "volume": ("Volume (m^3)",), "moments_of_inertia": ("Moments of inertia (te.m^2)",),
        "inertia_tensor_origin": ("Inertia tensor origin",),
        "radii_of_gyration_origin": ("Radii of gyration origin",),
    }
    result = {key: [] for key in names}
    for name in ("Summary", "Hydrostatics"):
        if name not in wb:
            continue
        sheet = wb[name]
        for row in sheet:
            key = next((key for key, labels in names.items() if row[0].value in labels), None)
            if key is None:
                continue
            count = 3 if key in ("radii_of_gyration", "moments_of_inertia") else 1
            rows = list(sheet.iter_rows(min_row=row[0].row, max_row=row[0].row + count - 1))
            values = [[c.value for c in r[1:] if c.value is not None] for r in rows]
            result[key].append(_statement(sheet, [c for r in rows for c in r], values=values))
    return {key: value or MISSING for key, value in result.items()}


def _read_conventions(wb):
    # Body/mesh/CG positions do not establish the reference point of the RAOs.
    names = {"Phase convention": "phase_convention",
             "Response reference point": "response_reference_point",
             "Heading convention": "heading_convention"}
    result = dict.fromkeys(names.values(), MISSING)
    for sheet in wb:
        for row in sheet:
            key = names.get(row[0].value)
            if key and len(row) > 1 and isinstance(row[1].value, str):
                result[key] = _statement(sheet, list(row[:2]), row[1].value)
    return result


def _source_metadata(path, fmt):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        return {"units": _read_units(wb, fmt), "source": _read_source(wb, fmt),
                "loading_condition": _read_loading(wb), "conventions": _read_conventions(wb)}
    finally:
        wb.close()


def _refresh_entry_metadata(entry):
    for field in METADATA_FIELDS:
        entry.metadata[field] = {source["path"]: source[field]
                                 for source in entry.metadata.get("sources", [])
                                 if field in source}


def _stored_basis(entry):
    basis = entry.metadata.get("basis", {})
    return {key: basis.get(key) for key in BASIS}


def _same_content(left, right) -> bool:
    return (left.hull_params == right.hull_params
            and _stored_basis(left) == _stored_basis(right)
            and all(np.array_equal(getattr(left.rao_data, field),
                                   getattr(right.rao_data, field)) for field in FIELDS))


def _merge_sources(retained, incoming) -> None:
    sources = retained.metadata.setdefault("sources", [])
    for source in incoming.metadata.get("sources", []):
        old = next((item for item in sources if item["path"] == source["path"]), None)
        if old is not None and old != source:
            raise ValueError("Changed provenance for an occupied source path")
        if old is None:
            sources.append(source)
    _refresh_entry_metadata(retained)


def store_fixture_entry(db: RAODatabase, entry, box_comparison=None):
    """Collapse exact content; permit only the separately verified box alias."""
    if any(not np.isfinite(getattr(entry.rao_data, field)).all() for field in FIELDS):
        raise ValueError("RAO arrays must be finite")
    for retained in db.query({}):
        exact = _same_content(retained, entry)
        box_alias = (box_comparison is not None and box_comparison["equivalent"]
                     and {retained.variation_id, entry.variation_id} == BOX_PAIR
                     and retained.hull_params == entry.hull_params
                     and _stored_basis(retained) == _stored_basis(entry))
        if retained.variation_id == entry.variation_id and not exact:
            raise ValueError("Different data for an occupied variation ID")
        if exact or box_alias:
            _merge_sources(retained, entry)
            if box_alias:
                retained.metadata["box_comparison"] = box_comparison
            return retained
    db.store(entry)
    return entry


def _source_provenance(root: Path, path: Path, retrieved_on: str) -> dict:
    relative = path.resolve().relative_to(root).as_posix()
    # Fixture root resolution must not inherit another checkout's Git binding.
    env = {key: value for key, value in os.environ.items()
           if key not in {"GIT_DIR", "GIT_WORK_TREE", "GIT_COMMON_DIR"}}
    revision = subprocess.run(
        ["git", "log", "-1", "--format=%H", "--", relative], cwd=root,
        env=env, check=True, capture_output=True, text=True,
    ).stdout.strip()
    if not revision:
        raise ValueError("Fixture has no source revision")
    raw = path.read_bytes()
    committed = subprocess.run(
        ["git", "show", f"{revision}:{relative}"], cwd=root,
        env=env, check=True, capture_output=True,
    ).stdout
    if raw != committed:
        raise ValueError("Fixture differs from its cited source revision")
    fmt = XlsxFormatDetector.detect(path)
    return {
        "path": relative, "sha256": hashlib.sha256(raw).hexdigest(),
        "revision": revision, "retrieved_on": retrieved_on, "format": fmt,
        "worksheet": "RAOs" if fmt == "pipeline" else "Displacement RAOs",
        "table": "all displacement RAO rows and headings",
        **_source_metadata(path, fmt),
        # Resolve from data/manifest on the repository's blob view.
        "url": f"../../{revision}/{relative}",
    }


def populate_fixture_database(db: RAODatabase, repo_root: Path, retrieved_on: str,
                              previous_sources=None):
    """Extract every named fixture, preserving aliases and repeat-run provenance."""
    root = Path(repo_root).resolve()
    date.fromisoformat(retrieved_on)
    previous = {source["path"]: source for entry in db.query({})
                for source in entry.metadata.get("sources", [])}
    previous.update(previous_sources or {})
    for stem in FIXTURE_STEMS:
        path = root / "tests/fixtures/solver" / f"{stem}.xlsx"
        source = _source_provenance(root, path, retrieved_on)
        old = previous.get(source["path"])
        if old is not None and old["sha256"] == source["sha256"]:
            source["retrieved_on"] = old["retrieved_on"]
        entry = populate_database_from_xlsx(
            RAODatabase(), path, stem, {},
            metadata={"source_file": source["path"],
                      "basis": BASIS.copy(), "sources": [source]},
        )
        _refresh_entry_metadata(entry)
        entry.metadata["stored_basis_qualification"] = (
            "Legacy storage declaration only; source units are recorded per source; "
            "unstated units are neither inferred nor converted")
        comparison = None
        if stem == "L00_test01":
            comparison = compare_box_exports(
                db.get_by_id("test01_unit_box").rao_data, entry.rao_data,
                root / "tests/fixtures/solver/test01_unit_box.xlsx", path,
            )
            if not comparison["equivalent"]:
                raise ValueError("Box alias comparison failed at source export precision")
        store_fixture_entry(db, entry, comparison)
    if {entry.variation_id for entry in db.query({})} != EXPECTED_IDS:
        raise ValueError("Unexpected fixture entry set; unrelated data cannot be replaced")
    return db


def _verify_readback(expected: RAODatabase, path: Path) -> None:
    restored = RAODatabase()
    restored.load_from_disk(path)
    if {entry.variation_id for entry in restored.query({})} != EXPECTED_IDS:
        raise ValueError("Parquet readback entry set differs")
    for entry in expected.query({}):
        saved = restored.get_by_id(entry.variation_id)
        if (not _same_content(entry, saved) or entry.metadata != saved.metadata
                or entry.rao_data.vessel_name != saved.rao_data.vessel_name):
            raise ValueError("Parquet readback content differs")


def regenerate_fixture_database(root: Path, output: Path, manifest: Path, retrieved_on: str):
    """Regenerate through repo population code; preserve divergent existing data."""
    output, manifest = Path(output), Path(manifest)
    if output.resolve() == manifest.resolve():
        raise ValueError("Database and manifest paths must differ")
    db = RAODatabase()
    if output.exists():
        db.load_from_disk(output)
        if {entry.variation_id for entry in db.query({})} - EXPECTED_IDS:
            raise ValueError("Existing database contains unrelated entries")
    existing = db
    previous = {s["path"]: s for e in existing.query({}) for s in e.metadata.get("sources", [])}
    db = populate_fixture_database(RAODatabase(), root, retrieved_on, previous)
    for old in existing.query({}):
        if not _same_content(old, db.get_by_id(old.variation_id)):
            raise ValueError("Different data for an occupied variation ID")
    source_date = db.get_by_id("test01_unit_box").metadata["sources"][0]["retrieved_on"]
    db.save_to_disk(output)
    _verify_readback(db, output)
    payload = {
        "owner": "digitalmodel", "dataset": "rao_database", "basis": BASIS,
        "basis_qualification": db.get_by_id("test01_unit_box").metadata["stored_basis_qualification"],
        "readiness": "benchmark fixtures; engineering coordinate/reference qualification unverified",
        "reuse": "Existing in-repository fixtures; repository LICENSE applies; no new upstream rights claimed",
        "regenerate": "PYTHONPATH=src uv run --no-sources python scripts/populate_rao_database.py "
                      f"--retrieved-on {source_date}",
        "comparison_criterion": "Named box pair only: per-source half decimal export unit (10**(order - observed significant digits + 1)/2) plus binary spacing; circular phase residual includes full-turn spacing. Generic deduplication requires exact arrays and matching basis.",
        "database_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
        "database_bytes": output.stat().st_size,
        "entries": [{"variation_id": entry.variation_id,
                     "frequency_count": len(entry.rao_data.frequencies),
                     "heading_count": len(entry.rao_data.directions),
                     "sources": entry.metadata["sources"],
                     "basis": entry.metadata["basis"],
                     **{field: entry.metadata[field] for field in METADATA_FIELDS},
                     **({"box_comparison": entry.metadata["box_comparison"]}
                        if "box_comparison" in entry.metadata else {})}
                    for entry in db.query({})],
    }
    manifest.parent.mkdir(parents=True, exist_ok=True)
    manifest.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    if json.loads(manifest.read_text(encoding="utf-8")) != payload:
        raise ValueError("Manifest readback differs")
    return db
