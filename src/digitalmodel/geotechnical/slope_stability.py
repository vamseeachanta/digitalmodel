# # Copyright (c) 2024 Digital Model Project
# # Licensed under the MIT License. See LICENSE file for details.

"""Optional xslope adapter for plane-strain slope and steady seepage screening."""

from __future__ import annotations

from dataclasses import dataclass, field
from importlib.metadata import version
from pathlib import Path
import tempfile
from typing import Any

from digitalmodel.geotechnical.slope_profile import _number, validate_profile


@dataclass(frozen=True)
class MethodResult:
    """One method and its own circular search convergence record."""

    factor_of_safety: float
    search_converged: bool | None
    critical_circle: dict[str, float] | None
    unsolved_trials: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class SeepageResult:
    """Discharge per unit out-of-plane width in the model length/time system."""

    discharge_per_width: float
    length_unit: str
    time_unit: str
    converged: bool
    closure_fraction: float


@dataclass(frozen=True)
class SlopeResult:
    """Screening results with the installed solver version and run settings."""

    xslope_version: str
    units: str
    num_slices: int
    methods: dict[str, MethodResult]
    seepage: SeepageResult | None = None


def _material_cells(layer: dict[str, Any], index: int, wet: bool) -> dict[str, Any]:
    cells = {
        "A": index + 1, "B": f"soil {index + 1}",
        "C": layer["unit_weight_kn_m3"],
        "D": layer.get("saturated_unit_weight_kn_m3"),
        "O": "piezo" if wet and "su_kpa" not in layer else "none",
        "AG": layer.get("permeability_m_s"),
        "AH": layer.get("permeability_m_s"), "AI": 0,
    }
    if "su_kpa" in layer:
        cells.update({"E": "cp", "F": layer["su_kpa"], "G": 0,
                      "H": layer["su_gradient_kpa_per_m"],
                      "I": layer["reference_elevation_m"]})
    else:
        cells.update({"E": "mc", "F": layer["c_kpa"], "G": layer["phi_deg"]})
    return cells


def _populate_workbook(workbook: Any, data: dict[str, Any]) -> None:
    main = workbook["main"]
    main["D8"], main["D9"] = "SI", "sec"
    main["D10"] = data["water_unit_weight_kn_m3"]
    main["D13"] = data["seismic_coefficient"]
    workbook["profile"]["B2"] = data["bottom_elevation_m"]
    for index, layer in enumerate(data["layers"]):
        for mat_column, value in _material_cells(layer, index, "water_table_m" in data).items():
            workbook["mat"][f"{mat_column}{11 + index}"] = value
        column = 1 + index * 3
        workbook["profile"].cell(5, column + 1, index + 1)
        for row, (x, y) in enumerate(layer["top_boundary_m"], 9):
            workbook["profile"].cell(row, column, x)
            workbook["profile"].cell(row, column + 1, y)
    for row, (x, y) in enumerate(data.get("water_table_m", []), 5):
        workbook["piezo"].cell(row, 1, x)
        workbook["piezo"].cell(row, 2, y)
    for row, circle in enumerate(data["circles"], 3):
        for column, key in [(2, "x_m"), (3, "y_m"), (8, "radius_m")]:
            workbook["circles"].cell(row, column, circle[key])
        workbook["circles"].cell(row, 4, "Radius")
    for index, head in enumerate(data.get("seepage", {}).get("heads_m", [])):
        column = 5 + index * 3
        workbook["seep bc"].cell(3, column + 1, head["head_m"])
        for row, (x, y) in enumerate(head["boundary_m"], 5):
            workbook["seep bc"].cell(row, column, x)
            workbook["seep bc"].cell(row, column + 1, y)


def build_slope_model(profile: dict[str, Any], scratch_directory: str | Path) -> dict[str, Any]:
    """Validate SI inputs and load a temporary copy of the installed template.

    Args:
        profile: Layered SI profile described in the slope screening guide.
        scratch_directory: Existing caller-owned directory for temporary workbooks.

    Returns:
        An xslope model containing geometry and soil properties, without outputs.
    """
    data = validate_profile(profile)
    from openpyxl import load_workbook
    from xslope.fileio import default_template_path, load_slope_data

    if version("xslope") != "1.0.3":
        raise ValueError("adapter requires benchmarked xslope version 1.0.3")
    with tempfile.TemporaryDirectory(dir=scratch_directory) as directory:
        workbook = load_workbook(default_template_path())
        _populate_workbook(workbook, data)
        path = Path(directory) / "profile.xlsx"
        workbook.save(path)
        workbook.close()
        model: dict[str, Any] = load_slope_data(path)
    return model


def run_seepage(model: dict[str, Any], element_size: float) -> SeepageResult:
    """Solve steady seepage and attach nodal pore pressures to the same model.

    Args:
        model: An xslope model with conductivities and head boundaries.
        element_size: Target linear-triangle mesh size in the model length unit.

    Returns:
        Discharge, units, convergence and relative mass-balance error.
    """
    if _number(element_size, "element_size", 0) == 0:
        raise ValueError("element_size must be positive")
    from xslope.mesh import build_mesh_from_polygons, build_polygons
    from xslope.seep import (
        apply_steady_stability_field, build_seep_data, run_seepage_analysis,
    )

    mesh = build_mesh_from_polygons(build_polygons(model), element_size, "tri3")
    model["mesh"] = mesh
    solution = run_seepage_analysis(build_seep_data(mesh, model))
    flow = _number(float(solution["flowrate"]), "discharge_per_width", 0)
    closure = _number(float(solution["closure_fraction"]), "closure_fraction", 0)
    if not solution["converged"]:
        raise RuntimeError("seepage did not converge; stability field is not accepted")
    apply_steady_stability_field(model, solution, verbose=False)
    for material in model["materials"]:
        if material["option"] != "cp":
            material["u"] = "seep"
    return SeepageResult(flow, "m" if model["unit_system"] == "si" else "ft",
                         model.get("time_unit") or "unspecified", True, closure)


def _solve_method(model: dict[str, Any], method: str, search: bool, slices: int) -> MethodResult:
    from xslope.search import circular_search
    from xslope.slice import generate_slices
    from xslope.solve import solve_selected

    unsolved: dict[str, Any] = {}
    if search:
        cache, converged, _, _ = circular_search(
            model, method, num_slices=slices, unsolved_out=unsolved
        )
        if not cache or not isinstance(cache[0].get("solver_result"), dict):
            raise RuntimeError(f"{method} search found no solved surface")
        best = cache[0]
        solved = best["solver_result"]
        if best["FS"] != solved.get("FS"):
            raise RuntimeError(f"{method} search returned a rejected surface")
        circle = {key: float(best[key]) for key in ("Xo", "Yo", "Depth")}
        circle["R"] = circle["Yo"] - circle["Depth"]
    else:
        success, generated = generate_slices(model, circle=model["circles"][0], num_slices=slices)
        if not success:
            raise RuntimeError(f"{method} slice generation failed")
        slice_df, _ = generated
        solved = solve_selected(method, slice_df)
        converged, circle = None, None
    if not solved or not solved.get("success", True):
        raise RuntimeError(f"{method} did not produce a solution")
    fs = _number(float(solved["FS"]), "factor_of_safety", 0)
    return MethodResult(fs, None if converged is None else bool(converged), circle, unsolved)


def analyze_model(
    model: dict[str, Any], *, methods: tuple[str, ...] = ("bishop", "spencer"),
    search: bool = True, num_slices: int = 20,
    seepage: SeepageResult | None = None,
) -> SlopeResult:
    """Run Bishop and/or Spencer on circular surfaces with explicit search status.

    Args:
        model: Model loaded by xslope or built from a validated SI profile.
        methods: Bishop and/or Spencer method names.
        search: Search from starting circles; false evaluates the first circle.
        num_slices: Requested slice count; xslope may split at layer boundaries.
        seepage: Optional record of the seepage solve already applied to the model.

    Returns:
        Versioned results without assigning a design acceptance criterion.
    """
    if type(num_slices) is not int or num_slices < 2:
        raise ValueError("num_slices must be an integer of at least two")
    if not methods or any(m not in ("bishop", "spencer") for m in methods):
        raise ValueError("methods must contain bishop and/or spencer")
    results = {m: _solve_method(model, m, search, num_slices) for m in methods}
    return SlopeResult(version("xslope"), model["unit_system"], num_slices, results, seepage)


def analyze_profile(
    profile: dict[str, Any], scratch_directory: str | Path, *, num_slices: int = 20,
) -> SlopeResult:
    """Run both circular searches, preceded by steady seepage when specified.

    Args:
        profile: Validated SI layered-profile input.
        scratch_directory: Existing directory for temporary solver input.
        num_slices: Requested number of slices for both methods.

    Returns:
        Bishop, Spencer and optional seepage results with the xslope version.
    """
    model = build_slope_model(profile, scratch_directory)
    seepage = None
    if "seepage" in profile:
        seepage = run_seepage(model, profile["seepage"]["element_size_m"])
    return analyze_model(model, num_slices=num_slices, seepage=seepage)
